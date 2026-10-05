#!/usr/bin/env python3
"""Phase 1 — Full DOCX paragraph export (MANUSCRIPT_WRITE=FALSE; read-only).

Writes manuscript-export/B1_P01-11_reader.paragraphs.jsonl
Each line: one python-docx body paragraph (index 0-based).
Also writes a w:p inventory sidecar for the expected 10944 / 10830 counts.
"""
from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "book1" / "source" / "B1_P01-11_reader.docx"
OUT_DIR = ROOT / "manuscript-export"
OUT_JSONL = OUT_DIR / "B1_P01-11_reader.paragraphs.jsonl"
OUT_META = OUT_DIR / "B1_P01-11_reader.export_meta.json"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

EXPECTED_SHA256 = "337c70ea468a587ab248f5eb49ffacc0c167cfd421a76d145c02dbc8c0595ecc"
EXPECTED_SIZE = 825302


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def wp_stats(path: Path) -> tuple[int, int]:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = etree.fromstring(xml)
    paras = root.xpath("//w:p", namespaces=NS)
    nonempty = 0
    for p in paras:
        joined = "".join((t.text or "") for t in p.xpath(".//w:t", namespaces=NS))
        if joined.strip():
            nonempty += 1
    return len(paras), nonempty


def flag_context(text: str) -> list[str]:
    flags: list[str] = []
    t = text or ""
    checks = [
        ("CITATION_BULLET", t.strip().startswith("•")),
        ("PAGE_CHECK_MARKER", "PAGE-CHECK" in t),
        ("SUTRA_MARKER", "সূত্র" in t),
        ("BARNITA_MARKER", "বর্ণিত আছে" in t),
        ("QURAN_MARKER", any(x in t for x in ("সূরা", "আয়াত", "কুরআন"))),
        ("SUNNI_BOOK", any(x in t for x in ("বুখারি", "মুসলিম", "তিরমিযি", "আবু দাউদ", "ইবনে মাজাহ", "নাসাঈ", "আহমদ", "হাকিম", "ত্ববারানী"))),
        ("SHIA_BOOK", any(x in t for x in ("কাফি", "বিহার", "ইহতিজাজ", "কুম্মী", "সাদূক", "সুল্লাইম", "আইয়াশী", "কুরবুল ইসনাদ"))),
        ("HADITH_WORD", "হাদীস" in t or "হাদিস" in t),
        ("RASUL_SAID", "রাসূল" in t and ("বলেছেন" in t or "বলেন" in t)),
        ("UMM_AYMAN", "উম্মে আইমান" in t or "উম্মেআইমান" in t),
        ("FADAK", "ফাদাক" in t),
        ("AUS", "আউস" in t or "আওস" in t),
    ]
    for name, ok in checks:
        if ok:
            flags.append(name)
    return flags


def main() -> None:
    if not DOCX.exists():
        sys.exit(f"missing DOCX: {DOCX}")
    digest = sha256_file(DOCX)
    size = DOCX.stat().st_size
    if digest != EXPECTED_SHA256 or size != EXPECTED_SIZE:
        sys.exit(f"DOCX lock FAIL sha256={digest} size={size}")

    wp_total, wp_nonempty = wp_stats(DOCX)
    doc = Document(str(DOCX))
    paras = list(doc.paragraphs)
    nonempty_body = sum(1 for p in paras if (p.text or "").strip())

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    n_flagged = 0
    with OUT_JSONL.open("w", encoding="utf-8") as out:
        for i, p in enumerate(paras):
            text = p.text or ""
            flags = flag_context(text)
            if flags:
                n_flagged += 1
            rec = {
                "para_index": i,
                "style": p.style.name if p.style is not None else "",
                "text": text,
                "nonempty": bool(text.strip()),
                "char_len": len(text),
                "flags": flags,
            }
            out.write(json.dumps(rec, ensure_ascii=False) + "\n")

    meta = {
        "docx_path": "book1/source/B1_P01-11_reader.docx",
        "sha256": digest,
        "size_bytes": size,
        "python_docx_paragraphs": len(paras),
        "python_docx_nonempty": nonempty_body,
        "w_p_total": wp_total,
        "w_p_nonempty": wp_nonempty,
        "flagged_paragraphs": n_flagged,
        "export_file": str(OUT_JSONL.relative_to(ROOT)),
        "note": "python-docx body paragraphs exclude nested table w:p; w_p_* counts all document.xml paragraphs.",
    }
    OUT_META.write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(meta, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
